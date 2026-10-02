#!/usr/bin/env python3
"""
to_canva.py — deck.json/deck.html -> deck.canva.pptx, built from NATIVE PowerPoint
shapes so Canva imports it as a FULLY EDITABLE design.

Why this exists (the Canva connector):
  Canva cannot ingest a local file over the MCP connector — `import-design-from-url`
  only accepts a already-public HTTPS URL, and client proposals must never be published
  to the open internet. Canva CAN import a local PPTX from Uploads, and a PPTX import
  keeps text boxes as real editable text. So the handoff is: build a PPTX made of real
  objects, then drop it into Canva.

  tools/build_pptx.py makes a FLAT pptx (one image per slide) that matches the PDF
  pixel-for-pixel — right for emailing, useless in Canva (nothing is editable).
  This tool is the opposite trade: every headline, rule, panel and photo is its own
  object, so Canva gets a real design. It sits ALONGSIDE build_pptx.py, it does not
  replace it.

How it works (two phases, one command):
  1. Playwright opens the rendered deck.html and reads the LIVE DOM — every box, fill,
     border, radius, background image, pseudo-element and text run — as flat drawing
     primitives in CSS px. Measuring beats hand-coded coordinates: the layout comes from
     CSS flex/grid, so reading it back means the PPTX lands where the PDF does and stays
     in sync whenever the templates change.
  2. python-pptx rebuilds those primitives as native shapes.

Output:  <proposal_dir>/<deck-name>.canva.pptx
Usage:   ./run.sh tools/to_canva.py proposals/<slug>/deck.json
         ./run.sh tools/to_canva.py proposals/<slug>              # dir also works
         ./run.sh tools/to_canva.py proposals/<slug>/other.html   # a second deck in the folder
         ./run.sh tools/to_canva.py proposals/<slug> --font-map 'CSS Name=Real Font'
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _util import PptxBuildError  # noqa: E402

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

EMU_PX = 9525          # 914400 EMU per inch / 96 px per inch
PT_PER_PX = 0.75

# ---------------------------------------------------------------------------
# PHASE 1 — the extractor. Runs inside the page, walks each .slide, and turns
# boxes, fills, radii, borders, images, pseudo-elements and text into primitives.
#
# Primitives (geometry in CSS px, origin = slide top-left):
#   {t:'rect'}  fills, hairline rules, borders, bullets   (+ radius -> shape kind)
#   {t:'grad'}  linear-gradient veils and scrims
#   {t:'img'}   <img> and url() backgrounds, crop resolved from object-fit
#   {t:'text'}  one text block, split into paragraphs and styled runs
# ---------------------------------------------------------------------------
EXTRACT_JS = r"""
(FONT_MAP) => {
  const OUT = [];
  const px = v => parseFloat(v) || 0;
  const num = v => { const n = parseFloat(v); return isNaN(n) ? null : n; };

  // "rgba(0, 0, 0, 0.5)" -> {hex:'000000', a:0.5}
  function parseColor(c) {
    if (!c) return null;
    const m = c.match(/rgba?\(([^)]+)\)/);
    if (!m) return null;
    const p = m[1].split(',').map(s => parseFloat(s));
    const a = p.length > 3 ? p[3] : 1;
    if (a === 0) return null;
    const hex = p.slice(0, 3).map(v => Math.round(v).toString(16).padStart(2, '0')).join('');
    return { hex: hex.toUpperCase(), a };
  }

  // linear-gradient(...) -> {angle, stops:[{pos,hex,a}]}
  function parseGradient(bg) {
    const i = bg.indexOf('linear-gradient(');
    if (i < 0) return null;
    let depth = 0, body = '', j = i + 'linear-gradient('.length;
    for (; j < bg.length; j++) {
      const ch = bg[j];
      if (ch === '(') depth++;
      else if (ch === ')') { if (depth === 0) break; depth--; }
      body += ch;
    }
    const parts = []; depth = 0; let cur = '';
    for (const ch of body) {
      if (ch === '(') depth++;
      if (ch === ')') depth--;
      if (ch === ',' && depth === 0) { parts.push(cur.trim()); cur = ''; }
      else cur += ch;
    }
    if (cur.trim()) parts.push(cur.trim());

    let angle = 180;                      // CSS default: to bottom
    const first = parts[0];
    if (/^(to |[\d.-]+deg)/.test(first)) {
      if (/deg/.test(first)) angle = parseFloat(first);
      else if (/to top/.test(first)) angle = 0;
      else if (/to bottom/.test(first)) angle = 180;
      else if (/to right/.test(first)) angle = 90;
      else if (/to left/.test(first)) angle = 270;
      parts.shift();
    }
    const stops = parts.map((p, k) => {
      const cm = p.match(/(rgba?\([^)]+\)|#[0-9a-f]+)/i);
      const col = parseColor(cm ? cm[1] : null) || { hex: '000000', a: 0 };
      const pm = p.match(/([\d.]+)%\s*$/);
      const pos = pm ? parseFloat(pm[1]) / 100 : (k === 0 ? 0 : k / (parts.length - 1));
      return { pos, hex: col.hex, a: col.a };
    });
    return { angle, stops };
  }

  // url("data:image/png;base64,...") -> the bare src
  function bgImageUrl(bgi) {
    if (!bgi || bgi === 'none' || bgi.includes('gradient(')) return null;
    const m = bgi.match(/url\((['"]?)(.*?)\1\)/);
    return m ? m[2] : null;
  }

  // border-radius -> how PowerPoint should draw the box.
  // A 50% radius is a dot, 16px is a card, 999px is a pill; all three are common here
  // and all three come out as ugly squares if the radius is dropped.
  function radiusOf(cs, w, h) {
    const r = px(cs.borderTopLeftRadius);
    if (r <= 0.5 || w <= 0 || h <= 0) return 0;
    return Math.min(r, Math.min(w, h) / 2);
  }

  // object-fit:cover / contain + object-position -> crop fractions off each edge
  function imageCrop(nw, nh, w, h, fit, ox, oy) {
    if (!nw || !nh || !w || !h) return { fit: 'fill', crop: [0, 0, 0, 0] };
    if (fit === 'contain') return { fit: 'contain', crop: [0, 0, 0, 0], nw, nh };
    if (fit !== 'cover') return { fit: 'fill', crop: [0, 0, 0, 0], nw, nh };
    const sa = nw / nh, ba = w / h;
    let l = 0, r = 0, t = 0, b = 0;
    if (sa > ba) { const c = 1 - ba / sa; l = c * ox; r = c * (1 - ox); }
    else { const c = 1 - sa / ba; t = c * oy; b = c * (1 - oy); }
    return { fit: 'cover', crop: [l, r, t, b], nw, nh };
  }

  // CSS family -> the typeface name written into the PPTX. Unmapped families pass
  // straight through, which is right when the CSS name IS the real font (Plus Jakarta
  // Sans, Space Mono, Archivo). Supply --font-map when a deck uses a CSS alias for a
  // differently-named file, e.g. a 'Client Display' alias for a differently named font file.
  function mapFont(ff) {
    const first = (ff || '').split(',')[0].replace(/["']/g, '').trim();
    return FONT_MAP[first] || first || 'Arial';
  }

  function runStyle(el, pseudo) {
    const cs = getComputedStyle(el, pseudo || null);
    const col = parseColor(cs.color) || { hex: '000000', a: 1 };
    return {
      font: mapFont(cs.fontFamily),
      size: px(cs.fontSize),
      weight: parseInt(cs.fontWeight) || 400,
      italic: cs.fontStyle === 'italic' || cs.fontStyle === 'oblique',
      color: col.hex,
      alpha: col.a,
      spacing: cs.letterSpacing === 'normal' ? 0 : px(cs.letterSpacing),
      caps: cs.textTransform === 'uppercase',
    };
  }

  // Collect styled runs from a text-bearing element. <br> starts a new paragraph.
  function collectRuns(el) {
    const paras = [[]];
    (function walk(node, style) {
      for (const child of node.childNodes) {
        if (child.nodeType === 3) {
          const raw = child.textContent.replace(/\s+/g, ' ');
          if (!raw.trim() && !raw.length) continue;
          const text = style.caps ? raw.toUpperCase() : raw;
          paras[paras.length - 1].push(Object.assign({}, style, { text }));
        } else if (child.nodeType === 1) {
          if (child.tagName === 'BR') { paras.push([]); continue; }
          if (child.tagName === 'IMG') continue;
          walk(child, runStyle(child));
        }
      }
    })(el, runStyle(el));
    return paras.map(rs => {
      const out = rs.filter(r => r.text.length);
      if (out.length) {
        out[0].text = out[0].text.replace(/^\s+/, '');
        out[out.length - 1].text = out[out.length - 1].text.replace(/\s+$/, '');
      }
      return out.filter(r => r.text.length);
    }).filter(rs => rs.length);
  }

  // A text leaf becomes ONE PowerPoint text box, so it must be a single flow of
  // inline text. Flex/grid parents lay their children out with gaps, and block
  // children stack — in both cases each child needs its own box, or the gaps
  // collapse and the strings run together ("PADEL RETREAT09:00 - 17:00").
  //
  // But display alone does not decide it: `.eyebrow` is inline-flex holding nothing
  // but a text node, and treating that as a container drops the text entirely (there
  // are no element children to recurse into). With no element children there are no
  // gaps to preserve, so it is always a leaf.
  function isTextLeaf(el) {
    if (!el.textContent.trim()) return false;
    const kids = [...el.children].filter(c => c.tagName !== 'BR');
    if (!kids.length) return true;
    if (/flex|grid/.test(getComputedStyle(el).display)) return false;
    for (const c of kids) {
      if (getComputedStyle(c).display !== 'inline') return false;
    }
    return true;
  }

  // The union of every text node's painted rect. A PowerPoint text box anchors its
  // content top-left, so a centring container (`place-items:center` on the avatar
  // chips, `align-items:center` on badges) would otherwise throw the glyphs into the
  // corner. Measuring where they actually land sidesteps the whole alignment problem.
  function textRectsUnion(el) {
    let box = null;
    (function walk(n) {
      for (const c of n.childNodes) {
        if (c.nodeType === 3 && c.textContent.trim()) {
          const rg = document.createRange();
          rg.selectNodeContents(c);
          const r = rg.getBoundingClientRect();
          if (!r.width && !r.height) continue;
          box = box ? { left: Math.min(box.left, r.left), top: Math.min(box.top, r.top),
                        right: Math.max(box.right, r.right),
                        bottom: Math.max(box.bottom, r.bottom) }
                    : { left: r.left, top: r.top, right: r.right, bottom: r.bottom };
        } else if (c.nodeType === 1 && c.tagName !== 'IMG') {
          walk(c);
        }
      }
    })(el);
    return box;
  }

  // Text sitting directly inside a container that also has element children would
  // otherwise be lost. Give each such text node its own box, measured where it paints.
  function orphanTextPrims(el, sr, opacity, prims) {
    const cs = getComputedStyle(el);
    for (const n of el.childNodes) {
      if (n.nodeType !== 3 || !n.textContent.trim()) continue;
      const rg = document.createRange();
      rg.selectNodeContents(n);
      const r = rg.getBoundingClientRect();
      if (!r.width && !r.height) continue;
      const st = runStyle(el);
      const lh = cs.lineHeight === 'normal' ? st.size * 1.2 : px(cs.lineHeight);
      const raw = n.textContent.replace(/\s+/g, ' ').trim();
      prims.push({ t: 'text', x: r.left - sr.left, y: r.top - sr.top,
                   w: r.width, h: r.height, lineHeight: lh, align: 'left', opacity,
                   paras: [[Object.assign({}, st,
                     { text: st.caps ? raw.toUpperCase() : raw })]] });
    }
  }

  // Where the text actually paints. Normally this is the content box, but an
  // inline-block ::before (the signature eyebrow dash) pushes the first line right;
  // the dash is not a DOM node, so only a Range knows where the text really starts.
  function firstTextRect(el) {
    for (const n of el.childNodes) {
      if (n.nodeType === 3 && n.textContent.trim()) {
        const rg = document.createRange();
        rg.selectNodeContents(n);
        const r = rg.getBoundingClientRect();
        if (r.width || r.height) return r;
      } else if (n.nodeType === 1 && n.textContent.trim()) {
        const r = firstTextRect(n);
        if (r) return r;
      }
    }
    return null;
  }

  // ::before / ::after. Four shapes appear in this deck: the full-bleed contour
  // texture, the inline eyebrow dash, absolutely-placed bullet chips, and the
  // connector rule between storyline dots.
  function pseudoPrims(el, which, x, y, w, h, opacity, prims, sr0) {
    const ps = getComputedStyle(el, which);
    if (!ps || !ps.content || ps.content === 'none' || ps.content === 'normal') return;

    const pw = px(ps.width), ph = px(ps.height);
    // In-flow vs absolutely placed. Test `position`, NOT `display`: the eyebrow dash is
    // declared inline-block, but as a flex item it gets blockified to `block`, so a
    // display test misreads it as absolute and drops it at the element origin.
    const inflow = ps.position === 'static';

    // geometry: absolutely placed pseudos carry real offsets; in-flow ones ride the
    // first line box, so centre them on the text we can measure.
    let bx = x, by = y, bw = pw || w, bh = ph || h;
    if (inflow) {
      const fr = firstTextRect(el);
      by = fr ? (fr.top - sr0.top) + (fr.height - bh) / 2 : y + (h - bh) / 2;
    } else {
      const l = num(ps.left), t = num(ps.top), b = num(ps.bottom), rt = num(ps.right);
      if (l !== null) bx = x + l; else if (rt !== null) bx = x + w - bw - rt;
      if (t !== null) by = y + t; else if (b !== null) by = y + h - bh - b;
    }

    // textual pseudo (the "›" list markers)
    const lit = ps.content.match(/^["'](.*)["']$/);
    if (lit && lit[1].trim()) {
      const st = runStyle(el, which);
      const lh = ps.lineHeight === 'normal' ? st.size * 1.2 : px(ps.lineHeight);
      prims.push({ t: 'text', x: bx, y: by, w: Math.max(bw, st.size * 1.6),
                   h: Math.max(bh, lh), lineHeight: lh, align: 'left',
                   opacity, paras: [[Object.assign({}, st,
                     { text: st.caps ? lit[1].toUpperCase() : lit[1] })]] });
      return;
    }

    if (bw <= 0 || bh <= 0) return;

    const src = bgImageUrl(ps.backgroundImage);
    if (src) {
      // mix-blend-mode has no PPTX equivalent; approximate the screen-blended
      // contour with straight alpha, which reads the same over a dark band.
      prims.push({ t: 'img', x: bx, y: by, w: bw, h: bh, src,
                   fit: ps.backgroundSize === 'contain' ? 'contain' : 'cover',
                   alpha: (parseFloat(ps.opacity) || 1) * opacity });
      return;
    }
    const grad = parseGradient(ps.backgroundImage);
    if (grad) {
      prims.push({ t: 'grad', x: bx, y: by, w: bw, h: bh,
                   angle: grad.angle, stops: grad.stops });
      return;
    }
    const bg = parseColor(ps.backgroundColor);
    if (bg) {
      prims.push({ t: 'rect', x: bx, y: by, w: bw, h: bh, fill: bg.hex,
                   alpha: bg.a * (parseFloat(ps.opacity) || 1) * opacity,
                   radius: radiusOf(ps, bw, bh) });
    }
  }

  document.querySelectorAll('.slide').forEach((slide, si) => {
    const sr = slide.getBoundingClientRect();
    const prims = [];
    const scs = getComputedStyle(slide);
    const sbg = parseColor(scs.backgroundColor);
    const out = { index: si, w: sr.width, h: sr.height,
                  bg: sbg ? sbg.hex : 'FFFFFF', prims };

    // `clip` carries down the nearest rounded overflow:hidden ancestor. Photos are not
    // rounded themselves — `.mood-img` (12px) and `.kol__avatar` (a 50% circle, so every
    // talent headshot) round them by clipping, and without this they import as squares.
    (function walk(el, clip) {
      const r = el.getBoundingClientRect();
      const x = r.left - sr.left, y = r.top - sr.top;
      const cs = getComputedStyle(el);
      if (cs.display === 'none' || cs.visibility === 'hidden') return;
      const opacity = parseFloat(cs.opacity);
      if (opacity === 0) return;
      const rad = radiusOf(cs, r.width, r.height);

      if (el !== slide) {
        // fill: colour first, then any gradient or image painted over it
        const bg = parseColor(cs.backgroundColor);
        if (bg) prims.push({ t: 'rect', x, y, w: r.width, h: r.height,
                             fill: bg.hex, alpha: bg.a * opacity, radius: rad });
        const bsrc = bgImageUrl(cs.backgroundImage);
        if (bsrc) {
          prims.push({ t: 'img', x, y, w: r.width, h: r.height, src: bsrc,
                       fit: cs.backgroundSize === 'contain' ? 'contain' : 'cover',
                       alpha: opacity });
        } else {
          const grad = parseGradient(cs.backgroundImage);
          if (grad) prims.push({ t: 'grad', x, y, w: r.width, h: r.height,
                                 angle: grad.angle, stops: grad.stops });
        }
        // borders, one thin rect per visible edge
        for (const [side, bx, by, bw, bh] of [
          ['Top', x, y, r.width, px(cs.borderTopWidth)],
          ['Bottom', x, y + r.height - px(cs.borderBottomWidth), r.width, px(cs.borderBottomWidth)],
          ['Left', x, y, px(cs.borderLeftWidth), r.height],
          ['Right', x + r.width - px(cs.borderRightWidth), y, px(cs.borderRightWidth), r.height],
        ]) {
          if (bh <= 0 || bw <= 0) continue;
          const bc = parseColor(cs['border' + side + 'Color']);
          if (!bc || cs['border' + side + 'Style'] === 'none') continue;
          prims.push({ t: 'rect', x: bx, y: by, w: bw, h: bh,
                       fill: bc.hex, alpha: bc.a * opacity, radius: 0 });
        }
      }
      pseudoPrims(el, '::before', x, y, r.width, r.height, opacity, prims, sr);

      if (el.tagName === 'IMG') {
        const ics = getComputedStyle(el);
        const opRaw = (ics.objectPosition || '50% 50%').split(/\s+/);
        const pos = v => v.endsWith('%') ? parseFloat(v) / 100 : 0.5;
        const ci = imageCrop(el.naturalWidth, el.naturalHeight, r.width, r.height,
                             ics.objectFit || 'fill', pos(opRaw[0]),
                             pos(opRaw[1] !== undefined ? opRaw[1] : '50%'));
        // filter and radius are how the brand system dresses a client logo: the
        // default cover whitens it with brightness(0) invert(1) so it reads on the
        // dark band, and logo_keep_color instead rounds it into a full-colour disc.
        // a clipping wrapper the image fills is the same thing as a radius on the image
        let irad = rad;
        if (!irad && clip && Math.abs(clip.x - x) < 1.5 && Math.abs(clip.y - y) < 1.5 &&
            Math.abs(clip.w - r.width) < 1.5 && Math.abs(clip.h - r.height) < 1.5) {
          irad = clip.r;
        }
        prims.push({ t: 'img', x, y, w: r.width, h: r.height,
                     src: el.getAttribute('src'), fit: ci.fit, crop: ci.crop,
                     nw: ci.nw, nh: ci.nh, alpha: opacity, radius: irad,
                     filter: cs.filter && cs.filter !== 'none' ? cs.filter : null });
        return;
      }

      if (isTextLeaf(el)) {
        const paras = collectRuns(el);
        if (paras.length) {
          const lh = cs.lineHeight === 'normal' ? px(cs.fontSize) * 1.2 : px(cs.lineHeight);
          let tx = x + px(cs.paddingLeft) + px(cs.borderLeftWidth);
          let tw = r.width - px(cs.paddingLeft) - px(cs.paddingRight)
                           - px(cs.borderLeftWidth) - px(cs.borderRightWidth);
          // an in-flow ::before (the eyebrow dash) indents the first line — start the
          // box where the glyphs actually are, or the text sits on top of the dash
          const bps = getComputedStyle(el, '::before');
          if (bps && bps.content && bps.content !== 'none' &&
              bps.content !== 'normal' && bps.position === 'static') {
            const fr = firstTextRect(el);
            if (fr) {
              const shift = (fr.left - sr.left) - tx;
              if (shift > 0.5) { tx += shift; tw -= shift; }
            }
          }
          let ty = y + px(cs.paddingTop) + px(cs.borderTopWidth);
          let th = r.height - px(cs.paddingTop) - px(cs.paddingBottom)
                            - px(cs.borderTopWidth) - px(cs.borderBottomWidth);
          // a flex/grid leaf centres its own text, so hug the measured glyphs instead
          // of the content box (+1px slack so PowerPoint's tighter tracking can't
          // force an extra wrap)
          if (/flex|grid/.test(cs.display)) {
            const u = textRectsUnion(el);
            if (u) {
              tx = u.left - sr.left;
              ty = u.top - sr.top;
              tw = (u.right - u.left) + 1;
              th = u.bottom - u.top;
            }
          }
          prims.push({
            t: 'text', x: tx, y: ty, w: tw, h: th,
            lineHeight: lh, align: cs.textAlign, opacity, paras,
          });
        }
        pseudoPrims(el, '::after', x, y, r.width, r.height, opacity, prims, sr);
        return;
      }

      // CSS paint order: flow content, then ::after, then z-raised children.
      orphanTextPrims(el, sr, opacity, prims);
      const sub = (/hidden|clip/.test(cs.overflow) && rad > 0)
        ? { x, y, w: r.width, h: r.height, r: rad } : clip;
      const kids = [...el.children];
      const raised = kids.filter(c => (parseInt(getComputedStyle(c).zIndex) || 0) > 0);
      const flow = kids.filter(c => !raised.includes(c));
      for (const c of flow) walk(c, sub);
      pseudoPrims(el, '::after', x, y, r.width, r.height, opacity, prims, sr);
      for (const c of raised) walk(c, sub);
    })(slide, null);

    OUT.push(out);
  });
  return OUT;
};
"""


def extract_layout(html_path: Path, font_map: dict[str, str] | None = None) -> list[dict]:
    """Open the rendered deck and read every slide back out of the live DOM."""
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1280, "height": 720})
        page.goto(html_path.resolve().as_uri(), wait_until="networkidle")
        page.wait_for_timeout(400)          # let the embedded fonts settle
        slides = page.evaluate(EXTRACT_JS, font_map or {})
        browser.close()
    if not slides:
        raise PptxBuildError(f"no .slide elements found in {html_path.name}")
    return slides


# ---------------------------------------------------------------------------
# PHASE 2 — rebuild the primitives as native PowerPoint shapes.
# ---------------------------------------------------------------------------
ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT,
         "center": PP_ALIGN.CENTER, "justify": PP_ALIGN.JUSTIFY,
         "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT}

# PowerPoint sets the first baseline from the font ascent; CSS centres the glyph in
# the line box. Nudging each text block up by this share of its leading lines them up.
BASELINE_NUDGE = 0.16

# A sub-pixel rectangle can fall between sample rows and never draw. Hairline rules and
# bullet chips are grown to this and re-centred — imperceptible, and they always paint.
MIN_THICKNESS = 1.5


def E(v: float) -> Emu:
    return Emu(int(round(v * EMU_PX)))


def sub(parent, tag: str):
    """Append a DrawingML child element."""
    from lxml import etree
    return etree.SubElement(parent, qn(tag))


def solid_fill_xml(el, hex_color: str, alpha: float = 1.0):
    fill = sub(el, "a:solidFill")
    clr = sub(fill, "a:srgbClr")
    clr.set("val", hex_color)
    if alpha < 0.999:
        sub(clr, "a:alpha").set("val", str(int(round(max(0.0, min(1.0, alpha)) * 100000))))
    return fill


def strip_fill_and_line(shape):
    """Drop the theme fill/line/effects an autoshape inherits, so we can set our own."""
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:blipFill", "a:pattFill"):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    for e in spPr.findall(qn("a:ln")):
        spPr.remove(e)
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def add_rect(slide, p):
    x, y, w, h = p["x"], p["y"], p["w"], p["h"]
    if w < MIN_THICKNESS:
        x -= (MIN_THICKNESS - w) / 2
        w = MIN_THICKNESS
    if h < MIN_THICKNESS:
        y -= (MIN_THICKNESS - h) / 2
        h = MIN_THICKNESS

    # radius decides the shape: a dot is an ellipse, a card is a rounded rect
    rad = p.get("radius", 0) or 0
    short = min(w, h)
    if rad >= short / 2 - 0.51:
        kind = MSO_SHAPE.OVAL
    elif rad > 0.5:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        kind = MSO_SHAPE.RECTANGLE

    sh = slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    sh.shadow.inherit = False
    if kind is MSO_SHAPE.ROUNDED_RECTANGLE:
        # python-pptx adjustment is the corner radius as a fraction of the short side
        sh.adjustments[0] = max(0.0, min(0.5, rad / short))
    strip_fill_and_line(sh)
    solid_fill_xml(sh._element.spPr, p["fill"], p.get("alpha", 1.0))
    sub(sub(sh._element.spPr, "a:ln"), "a:noFill")
    sh.text_frame.word_wrap = False
    return sh


def add_grad(slide, p):
    """Translucent veils and scrims, baked to a 32-bit PNG.

    DrawingML can express per-stop alpha, but renderers disagree about honouring it
    (macOS QuickLook drops it and paints the veil solid black). A PNG is unambiguous
    everywhere, and it is still its own object — selectable, movable, deletable.
    """
    w, h = max(int(round(p["w"])), 1), max(int(round(p["h"])), 1)
    stops = sorted(p["stops"], key=lambda s: s["pos"])
    ang = p["angle"] % 360

    n = max(h if ang in (0, 180) else w, 2)
    vertical = ang in (0, 180)
    strip = Image.new("RGBA", (1, n) if vertical else (n, 1))
    px = strip.load()
    for i in range(n):
        t = i / (n - 1)
        if ang in (0, 270):                 # "to top" / "to left": stop 0 sits at the end
            t = 1 - t
        lo, hi = stops[0], stops[-1]
        for a, b in zip(stops, stops[1:]):
            if a["pos"] <= t <= b["pos"]:
                lo, hi = a, b
                break
        span = (hi["pos"] - lo["pos"]) or 1
        k = max(0.0, min(1.0, (t - lo["pos"]) / span))
        rgb_lo = tuple(int(lo["hex"][j:j + 2], 16) for j in (0, 2, 4))
        rgb_hi = tuple(int(hi["hex"][j:j + 2], 16) for j in (0, 2, 4))
        col = tuple(int(round(rgb_lo[c] + (rgb_hi[c] - rgb_lo[c]) * k)) for c in range(3))
        alpha = int(round((lo["a"] + (hi["a"] - lo["a"]) * k) * 255))
        px[(0, i) if vertical else (i, 0)] = col + (alpha,)

    buf = io.BytesIO()
    strip.resize((w, h), Image.BILINEAR).save(buf, "PNG", optimize=True)
    buf.seek(0)
    return slide.shapes.add_picture(buf, E(p["x"]), E(p["y"]), E(p["w"]), E(p["h"]))


DATA_URI = re.compile(r"^data:image/([\w+.-]+);base64,(.*)$", re.S)


def decode_img(src: str) -> io.BytesIO:
    m = DATA_URI.match(src or "")
    if not m:
        raise PptxBuildError(
            f"image src is not an inlined data URI: {(src or '')[:60]!r} — "
            "run build_deck.py so assets are embedded")
    return io.BytesIO(base64.b64decode(m.group(2)))


CSS_FILTER = re.compile(r"(\w[\w-]*)\(([^)]*)\)")


def _filter_amount(raw: str, default: float = 1.0) -> float:
    raw = raw.strip()
    if not raw:
        return default
    if raw.endswith("%"):
        return float(raw[:-1]) / 100
    try:
        return float(raw)
    except ValueError:
        return default


def _apply_css_filter(im: "Image.Image", css: str) -> "Image.Image":
    """Apply the filter functions the deck actually uses, in declaration order.

    `brightness(0) invert(1)` is the brand system's way of flattening a client logo to
    a white silhouette for a dark band; without it the logo keeps its original colours
    and can disappear against the near-black.
    """
    from PIL import ImageOps

    rgb, a = im.convert("RGBA").split()[:3], im.getchannel("A")
    im = Image.merge("RGB", rgb)
    for name, arg in CSS_FILTER.findall(css):
        n = name.lower()
        if n == "brightness":
            k = _filter_amount(arg)
            im = im.point(lambda v, k=k: max(0, min(255, int(round(v * k)))))
        elif n == "invert":
            k = _filter_amount(arg)
            im = Image.blend(im, ImageOps.invert(im), max(0.0, min(1.0, k)))
        elif n == "grayscale":
            k = _filter_amount(arg)
            im = Image.blend(im, ImageOps.grayscale(im).convert("RGB"),
                             max(0.0, min(1.0, k)))
        elif n == "opacity":
            k = max(0.0, min(1.0, _filter_amount(arg)))
            a = a.point(lambda v, k=k: int(round(v * k)))
        # anything else (blur, drop-shadow, ...) is left alone rather than guessed at
    im = im.convert("RGBA")
    im.putalpha(a)
    return im


def _corner_mask(size: tuple[int, int], radius: float) -> "Image.Image":
    """An alpha mask with rounded corners — a full-radius mask is a circle/ellipse."""
    from PIL import ImageDraw

    w, h = size
    mask = Image.new("L", size, 0)
    d = ImageDraw.Draw(mask)
    r = max(0.0, min(radius, min(w, h) / 2))
    if r >= min(w, h) / 2 - 0.51:
        d.ellipse((0, 0, w - 1, h - 1), fill=255)
    else:
        d.rounded_rectangle((0, 0, w - 1, h - 1), radius=r, fill=255)
    return mask


def add_img(slide, p):
    stream = decode_img(p["src"])
    x, y, w, h = p["x"], p["y"], p["w"], p["h"]
    fit = p.get("fit", "fill")
    crop = p.get("crop")
    nw, nh = p.get("nw"), p.get("nh")

    # url() backgrounds have no naturalWidth in the DOM — measure the decoded bytes
    if (nw is None or nh is None) and fit in ("cover", "contain"):
        with Image.open(stream) as im:
            nw, nh = im.size
        stream.seek(0)

    if fit == "contain" and nw and nh:
        sa, ba = nw / nh, w / h            # letterbox exactly like object-fit:contain
        if sa > ba:
            nh2 = w / sa
            y += (h - nh2) / 2
            h = nh2
        else:
            nw2 = h * sa
            x += (w - nw2) / 2
            w = nw2

    if fit == "cover" and crop is None and nw and nh:
        sa, ba = nw / nh, w / h            # centred cover crop
        if sa > ba:
            c = 1 - ba / sa
            crop = [c / 2, c / 2, 0, 0]
        else:
            c = 1 - sa / ba
            crop = [0, 0, c / 2, c / 2]

    # Everything below has to live in the PIXELS, not in DrawingML. PPTX can express
    # picture transparency (<a:alphaModFix>), but renderers drop it — macOS QuickLook
    # ignores it outright — and a texture meant to sit at 24% then paints as an opaque
    # sheet over the dark band and the white headline on top of it. CSS filters and
    # rounded corners have no PPTX equivalent at all. Bake all three into the bitmap.
    alpha = p.get("alpha", 1.0)
    filt = p.get("filter")
    rad = p.get("radius", 0) or 0

    if alpha < 0.999 or filt or rad > 0.5:
        im = Image.open(stream).convert("RGBA")

        # A rounded picture must be masked at its FINAL shape, so resolve any cover
        # crop into real pixels first, then mask — otherwise the corners get stretched.
        if rad > 0.5:
            if crop:
                W, H = im.size
                l, r_, t, b = crop
                im = im.crop((round(l * W), round(t * H),
                              max(round(W - r_ * W), round(l * W) + 1),
                              max(round(H - b * H), round(t * H) + 1)))
                crop = None
            ss = 3          # supersample so the curve stays clean at slide scale
            im = im.resize((max(int(w * ss), 1), max(int(h * ss), 1)), Image.LANCZOS)
            im.putalpha(_corner_mask(im.size, rad * ss))

        if filt:
            im = _apply_css_filter(im, filt)
        if alpha < 0.999:
            im.putalpha(im.getchannel("A").point(lambda v: int(round(v * alpha))))

        baked = io.BytesIO()
        im.save(baked, "PNG", optimize=True)
        baked.seek(0)
        stream = baked

    pic = slide.shapes.add_picture(stream, E(x), E(y), E(w), E(h))
    if fit == "cover" and crop:
        pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = crop
    return pic


def is_emoji(text: str) -> bool:
    """True when every visible character sits in an emoji/pictograph block."""
    chars = [c for c in text if not c.isspace()]
    return bool(chars) and all(
        0x1F000 <= ord(c) <= 0x1FAFF or 0x2600 <= ord(c) <= 0x27BF or ord(c) == 0xFE0F
        for c in chars)


def add_text(slide, p):
    lead = p["lineHeight"]
    # Exact CSS content width: PowerPoint tracks slightly tighter than the browser, so
    # any extra here lets a line that wrapped in the PDF fit on one line, which leaves
    # the measured rules and dividers sitting a line too low.
    box = slide.shapes.add_textbox(
        E(p["x"]), E(p["y"] - lead * BASELINE_NUDGE),
        E(p["w"]), E(max(p["h"], lead) + lead))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit"):
        for e in bodyPr.findall(qn(tag)):
            bodyPr.remove(e)
    sub(bodyPr, "a:noAutofit")

    opacity = p.get("opacity", 1.0)
    for i, runs in enumerate(p["paras"]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = ALIGN.get(p.get("align", "left"), PP_ALIGN.LEFT)
        para.line_spacing = Pt(lead * PT_PER_PX)
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        for rs in runs:
            run = para.add_run()
            run.text = rs["text"]
            f = run.font
            f.name = rs["font"]
            f.size = Pt(round(rs["size"] * PT_PER_PX, 1))
            # PPTX has no variable-font axis: the 800-weight display face maps to bold.
            f.bold = rs["weight"] >= 600
            f.italic = bool(rs["italic"])
            rPr = run._r.get_or_add_rPr()
            # Set the colour through python-pptx so <a:solidFill> is INSERTED in schema
            # order. Appending it by hand lands it after <a:latin>, which is invalid
            # DrawingML — PowerPoint tolerates it, a stricter importer can drop the
            # colour and render the run black.
            f.color.rgb = RGBColor.from_string(rs["color"])
            run_alpha = rs.get("alpha", 1.0) * opacity
            if run_alpha < 0.999:
                srgb = rPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                sub(srgb, "a:alpha").set(
                    "val", str(int(round(max(0.0, min(1.0, run_alpha)) * 100000))))
            if rs.get("spacing"):
                rPr.set("spc", str(int(round(rs["spacing"] * PT_PER_PX * 100))))
            # Emoji live in no brand face — leave the typeface unpinned so the system
            # emoji font is used instead of rendering tofu.
            if not is_emoji(rs["text"]):
                for tag in ("a:latin", "a:ea", "a:cs"):
                    el = rPr.find(qn(tag))
                    if el is None:
                        el = sub(rPr, tag)
                    el.set("typeface", rs["font"])
            else:
                f.name = None
    return box


def merge_rules(prims: list[dict]) -> list[dict]:
    """Join collinear, touching, identically-filled rectangles into single shapes.

    Every table cell contributes its own border-bottom, so one calendar can emit ~45
    separate hairlines where the design has 9 rules. That is a lot of loose objects to
    click past in the Canva editor, and abutting shapes can show antialiasing seams.
    """
    LOOKBACK = 12          # bounded so merging can never reorder distant shapes
    out: list[dict] = []
    for p in prims:
        merged = False
        if p["t"] == "rect" and not (p.get("radius") or 0) > 0.5:
            for q in reversed(out[-LOOKBACK:]):
                if q["t"] != "rect" or (q.get("radius") or 0) > 0.5:
                    continue
                if (q["fill"] != p["fill"]
                        or abs(q.get("alpha", 1) - p.get("alpha", 1)) > 1e-6):
                    continue
                if (abs(q["y"] - p["y"]) < 0.05 and abs(q["h"] - p["h"]) < 0.05
                        and abs(q["x"] + q["w"] - p["x"]) < 0.6):
                    q["w"] = p["x"] + p["w"] - q["x"]      # extend rightwards
                    merged = True
                    break
                if (abs(q["x"] - p["x"]) < 0.05 and abs(q["w"] - p["w"]) < 0.05
                        and abs(q["y"] + q["h"] - p["y"]) < 0.6):
                    q["h"] = p["y"] + p["h"] - q["y"]      # extend downwards
                    merged = True
                    break
        if not merged:
            out.append(p)
    return out


def build_pptx(slides: list[dict], out_path: Path) -> dict:
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(slides[0]["w"]), E(slides[0]["h"])
    blank = prs.slide_layouts[6]

    counts = {"rect": 0, "grad": 0, "img": 0, "text": 0}
    for s in slides:
        sl = prs.slides.add_slide(blank)
        # Set the real page background AND paint a full-bleed rectangle over it.
        # p:bg is the correct PPTX way, but not every consumer honours it (macOS
        # QuickLook ignores it outright), and a dropped background turns every dark
        # slide into white-on-white. The rectangle renders everywhere.
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor.from_string(s["bg"])
        add_rect(sl, {"x": 0, "y": 0, "w": s["w"], "h": s["h"],
                      "fill": s["bg"], "alpha": 1.0, "radius": 0})
        for p in merge_rules(s["prims"]):
            if p["t"] == "rect":
                add_rect(sl, p)
            elif p["t"] == "grad":
                add_grad(sl, p)
            elif p["t"] == "img":
                add_img(sl, p)
            elif p["t"] == "text":
                add_text(sl, p)
            counts[p["t"]] += 1

    prs.save(str(out_path))
    return counts


# ---------------------------------------------------------------------------
def resolve_deck(target: Path) -> tuple[Path, Path]:
    """Accept a proposal dir, a deck.json, or ANY rendered .html; return (dir, html).

    Naming it explicitly matters: a proposal folder can hold more than one deck (the
    folder may carry both a partnership deck and an event deck), so an
    .html argument is honoured as-is rather than being rewritten to deck.html.
    """
    if target.is_dir():
        proposal_dir = target
    elif target.suffix in (".json", ".html"):
        proposal_dir = target.parent
    else:
        raise PptxBuildError(f"expected a proposal dir, deck.json or a .html deck — got {target}")

    if target.suffix == ".html":
        if not target.exists():
            raise PptxBuildError(f"{target} does not exist")
        return proposal_dir, target

    html_path = proposal_dir / "deck.html"
    if not html_path.exists():
        deck_json = proposal_dir / "deck.json"
        if not deck_json.exists():
            raise PptxBuildError(f"no deck.html or deck.json in {proposal_dir}")
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from build_deck import render_html  # noqa: PLC0415
        render_html(deck_json)
    return proposal_dir, html_path


def main() -> int:
    ap = argparse.ArgumentParser(description="Build a Canva-ready editable PPTX from a deck.")
    ap.add_argument("target", type=Path,
                    help="proposals/<slug>/ (or its deck.json / deck.html)")
    ap.add_argument("--out", type=Path, default=None,
                    help="output path (default: <proposal_dir>/<deck-name>.canva.pptx)")
    ap.add_argument("--layout", action="store_true",
                    help="also write layout.json (the measured primitives, for debugging)")
    ap.add_argument("--font-map", action="append", default=[], metavar="CSS=PPTX",
                    help="rename a CSS font family in the output, e.g. "
                         "--font-map 'Client Display=RealFontName' (repeatable)")
    args = ap.parse_args()

    font_map = {}
    for pair in args.font_map:
        if "=" not in pair:
            print(f"ERROR: --font-map expects CSS=PPTX, got {pair!r}", file=sys.stderr)
            return 1
        css, _, pptx_name = pair.partition("=")
        font_map[css.strip()] = pptx_name.strip()

    try:
        proposal_dir, html_path = resolve_deck(args.target.resolve())
        slides = extract_layout(html_path, font_map)
        if args.layout:
            (proposal_dir / "layout.json").write_text(json.dumps(slides), encoding="utf-8")
        # deck.html -> deck.canva.pptx, event-deck.html -> event-deck.canva.pptx
        out = args.out or proposal_dir / f"{html_path.stem}.canva.pptx"
        counts = build_pptx(slides, out)
    except PptxBuildError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1

    print(f"Built {out.relative_to(Path.cwd()) if out.is_relative_to(Path.cwd()) else out} "
          f"({out.stat().st_size // 1024} KB) — {len(slides)} slides, fully editable: "
          f"{counts['text']} text boxes, {counts['img']} pictures, "
          f"{counts['rect']} shapes, {counts['grad']} gradients")
    print("Next: upload it to Canva (Uploads → it imports as an editable presentation).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
