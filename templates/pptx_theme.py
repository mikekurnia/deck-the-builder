"""
pptx_theme.py — brand constants + helpers for the editable PPTX renderer.
Mirrors tokens.css so the .pptx reads as the agency brand even though it's not pixel-identical
to the HTML/PDF. Fonts are named (Plus Jakarta Sans / Space Mono); PowerPoint
substitutes if the machine lacks them.
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# 16:9 canvas
EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

# Brand colours (from tokens.css)
INK = RGBColor(0x0E, 0x15, 0x12)
PRIMARY = RGBColor(0x1F, 0x52, 0x40)
PRIMARY_LIGHT = RGBColor(0xB7, 0xD3, 0xC4)
ACCENT = RGBColor(0xC9, 0x9A, 0x45)
ACCENT_LIGHT = RGBColor(0xDD, 0xB6, 0x6C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY = RGBColor(0x30, 0x3A, 0x49)
MUTED = RGBColor(0x69, 0x75, 0x86)
BORDER = RGBColor(0xE3, 0xE8, 0xEF)
GREEN = RGBColor(0x18, 0xA9, 0x57)
BG_SUBTLE = RGBColor(0xF8, 0xFA, 0xFC)

FONT_DISPLAY = "Plus Jakarta Sans"
FONT_MONO = "Space Mono"

MARGIN = Inches(0.9)
CONTENT_W = EMU_W - MARGIN * 2


def set_bg(slide, color):
    """Solid-fill a slide background."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_box(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def style_run(run, size, color, bold=False, font=FONT_DISPLAY, spacing=None, caps=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    if caps:
        run.text = run.text.upper()


def eyebrow(tf, text, on_dark=False):
    """Space Mono uppercase eyebrow line into the first paragraph of tf."""
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = (text or "").upper()
    style_run(r, 11, ACCENT_LIGHT if on_dark else PRIMARY, bold=True, font=FONT_MONO)


def para(tf, text, size, color, bold=False, font=FONT_DISPLAY, space_before=6, new=True):
    p = tf.add_paragraph() if new else tf.paragraphs[0]
    p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    style_run(r, size, color, bold=bold, font=font)
    return p


def bullets(tf, items, size=13, color=BODY):
    for it in items:
        p = tf.add_paragraph()
        p.space_before = Pt(5)
        r = p.add_run()
        r.text = "›  " + it
        style_run(r, size, color, font=FONT_DISPLAY)
